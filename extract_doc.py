import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
doc = Document(r'C:\Users\surya\OneDrive\Desktop\final_plan_revised.docx')
for p in doc.paragraphs:
    print(p.text)

# Also extract table content
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            print(cell.text)
